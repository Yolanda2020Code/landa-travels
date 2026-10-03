"""Create blank, practical user-study materials. Never creates participant results."""
import csv
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parent / "user-testing-pack"
TASKS = [
    ("T1", "Plan a short city break",
     "Plan a Berlin–Paris trip for two adults, using dates 4–6 weeks in the future, "
     "a total budget of EUR 1,000 and a balanced sustainability priority. Choose dates before the sessions and keep them the same for every participant. "
     "Use the controls or text naturally. Review the summary, approve it and inspect recommendations.",
     "The review shows the intended route, dates, party, currency/budget and priority. "
     "After approval, the participant reaches results or a clearly explained unavailable outcome. "
     "Record reaching results separately: an honest unavailable response is not proof of available inventory.", "10 minutes"),
    ("T2", "Correct one detail",
     "Change only the departure city from Berlin to Leipzig. Check the other details and approve the revised plan when asked.",
     "Origin changes; Paris, party and other unchanged details remain correct. "
     "The participant recognises that changed results require the revised context/approval.", "5 minutes"),
    ("T3", "Explain a carbon recommendation",
     "Find a carbon estimate or comparison and tell the observer what it means, which source/basis is shown, "
     "and one reason it might be incomplete. Explain whether a mapped or certified hotel guarantees a room or price.",
     "Participant distinguishes estimate from guarantee and can identify source/basis or explicitly reports that it cannot be found. "
     "Score comprehension from their answer, not merely whether a button opens.", "5 minutes"),
    ("T4", "Review the human-help path",
     "Ask for a human advisor. Inspect the summary and consent control and explain what would be shared and what happens next. "
     "Do not enter real contact details or submit a live handover for this study.",
     "The participant finds the path, understands consent and the bounded context, and recognises asynchronous replies. "
     "This task tests discovery/comprehension, not actual submission, advisor delivery or receipt.", "5 minutes"),
]
QUESTIONS = [
    ("Q1", "The assistant's questions and instructions were clear."),
    ("Q2", "It was easy to complete the planning tasks."),
    ("Q3", "I understood why a travel option was recommended."),
    ("Q4", "I could understand the sources and limitations of carbon or sustainability claims."),
    ("Q5", "I felt in control when correcting details or deciding whether to seek human help."),
    ("Q6", "Based on the labels and limitations shown, I could judge when to rely on the assistant's information."),
]


def table(d, headers, rows):
    t = d.add_table(rows=1, cols=len(headers))
    t.style = "Light Shading Accent 1"
    for c, text in zip(t.rows[0].cells, headers):
        c.text = text
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for row in rows:
        for c, text in zip(t.add_row().cells, row):
            c.text = text
    for row in t.rows:
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)
    return t


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    d = Document()
    for s in d.sections:
        s.top_margin = s.bottom_margin = Inches(.65)
        s.left_margin = s.right_margin = Inches(.7)
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(10)
    d.styles["Normal"].paragraph_format.space_after = Pt(6)
    for name in ["Heading 1", "Heading 2"]:
        d.styles[name].font.color.rgb = RGBColor.from_string("235A46")
    d.add_heading("Landa Travels: real-user testing pack", 0)
    d.add_paragraph("Blank protocol and forms for 3–5 adult volunteers. No participant results are included.")
    d.add_heading("1. Recruit and prepare", 1)
    d.add_paragraph(
        "Invite 3–5 adult friends, classmates or potential travellers who did not build the chatbot. "
        "Include a mix of travel/digital experience where practical. This is a small convenience sample, "
        "not representative evidence of usability, accessibility, trust or behaviour change.")
    d.add_paragraph(
        "Allow 25–35 minutes per person. Use the same public app version and task sheet, test people individually, "
        "and note date/time, source revision if known, browser, device and any connection problems. "
        "Use synthetic trip details and participant IDs P01–P05, not names. Do not test while local cross-validation is running.")
    d.add_paragraph(
        "Public application: https://yolandankala-landa-travels.hf.space. Prefer a direct browser tab, "
        "not an embedded frame. Do not require account login, payments, real CVs or live advisor submissions. "
        "Technical blocks must be recorded rather than worked around and scored as success.")
    d.add_heading("2. Consent and think-aloud script", 1)
    d.add_paragraph(
        "Read: 'I am evaluating the website, not you. Participation is voluntary; you can skip a question or stop at any time. "
        "Please use only the fictional trip details. I will record anonymous notes about actions, difficulties and your survey answers. "
        "I will not record your screen, voice, name or contact details unless we agree separately. "
        "The notes may be summarised anonymously in my assignment. Is that okay?'")
    d.add_paragraph(
        "Record consent yes/no locally. Agree a deletion date with the participant and follow institutional requirements. "
        "Do not proceed after a refusal. Any recording or verbatim quotation needs separate permission.")
    d.add_paragraph(
        "Read: 'As you use the site, say what you are looking for, what you expect to happen and anything confusing. "
        "You can say “I expected…” or “I cannot find…”. I will mostly observe without helping.' "
        "After silence, ask only 'What are you thinking now?' Do not suggest the correct control or desired opinion.")
    d.add_heading("3. Task sheet and success criteria", 1)
    for task_id, title, instructions, success, limit in TASKS:
        d.add_heading(f"{task_id}. {title}", 2)
        d.add_paragraph("Participant instruction: " + instructions)
        d.add_paragraph("Observer criterion: " + success)
        d.add_paragraph("Time cap: " + limit + ". Start timing when the task is read; stop at the outcome or cap.")
    d.add_heading("4. Observer recording sheet", 1)
    d.add_paragraph(
        "For each task record: participant ID, outcome (independent success / assisted success / partial / failed / "
        "technical block / not attempted), elapsed seconds, help count, errors, observed difficulty and an anonymised note. "
        "If the participant is stuck, let them stop or record help explicitly. Never silently reset and erase a failed attempt.")
    table(d, ["ID / task", "Outcome", "Seconds / help", "Observation"], [
        ("____ / T1", "", "", ""), ("____ / T2", "", "", ""),
        ("____ / T3", "", "", ""), ("____ / T4", "", "", ""),
    ])
    d.add_heading("5. Post-task five-point survey", 1)
    d.add_paragraph(
        "Ask immediately after the tasks, without influencing answers. "
        "1 = strongly disagree; 2 = disagree; 3 = neither agree nor disagree; 4 = agree; 5 = strongly agree. "
        "N/A = not experienced or unable to judge. These are study-specific items, not a validated SUS or trust scale.")
    table(d, ["Item", "Statement", "Rating 1–5 / N/A"],
          [(key, question, "") for key, question in QUESTIONS])
    d.add_paragraph(
        "Open questions: What was most confusing? What helped you understand a recommendation? "
        "What one change would make this easier? Would you need to check any information elsewhere before booking?")
    d.add_heading("6. Analyse honestly", 1)
    d.add_paragraph(
        "Count independent and assisted success separately for each task. Report technical blocks and not-attempted tasks "
        "with denominators; do not turn missing tasks into successful ones. Summarise elapsed times as medians and ranges, "
        "with capped/failed attempts identified. Count recurring issues and prioritise them by impact.")
    d.add_paragraph(
        "For each survey item report the number of valid responses and median (optionally range). "
        "Leave N/A out of that item's score but show its count. With 3–5 people, do not run significance claims "
        "or conclude general usability, environmental behaviour change, GDPR compliance or screen-reader accessibility.")
    d.add_paragraph(
        "Keep earlier automated proxy scores in a separate table. They are not human Likert responses. "
        "If you later fix an issue, report it as an observed problem and proposed/implemented response; "
        "claim improvement only if a genuine follow-up test measured it.")
    d.add_heading("7. What to add to your report", 1)
    d.add_paragraph(
        "Replace or extend Section 6.4 with the actual study method, participant count, outcomes and limitations. "
        "Add one task-results table (success/assistance/blocks/time) and one survey table (n/median/range), "
        "then 2–3 concrete findings linked to design changes. An anonymised observation screenshot may help if separately consented. "
        "Append the task sheet and questionnaire. Keep the main-text limit and institutional counting rules in mind.")
    d.add_paragraph(
        "Template only—fill every bracket after real testing: "
        "'A convenience sample of [n] adult volunteers completed four planning/comprehension tasks using [version/date] "
        "on [devices]. Sessions used think-aloud observation followed by six five-point items. "
        "Independent task success was [per-task counts], with [technical blocks/assistance]. "
        "The main difficulties were [actual observations]. Survey item medians were [actual scores]. "
        "The small convenience sample and [specific limitations] prevent generalisation; earlier proxy scores remain separate.'")
    d.add_paragraph(
        "Never fill blanks with expected results. A completed study does not guarantee a particular mark.")
    d.add_heading("8. Highest-value evidence additions to this report", 1)
    d.add_paragraph(
        "Your uploaded report already has extensive architecture, dialogue, data, privacy and deployment visuals. "
        "Do not add another general architecture diagram. Prioritise (1) actual participant task/survey tables in Section 6.4; "
        "(2) a five-fold metrics table and aggregate confusion matrix in Section 6.2 after the run finishes; "
        "(3) a short real provider-error-handling excerpt in Section 5.3, beside the existing ranking excerpt. "
        "Use the confusion matrix to discuss actual intent confusions, not just display a good overall score.")
    d.add_paragraph(
        "Cross-validation trains five separate NLU models on four folds each and tests on the remaining fold; "
        "it does not evaluate or replace the existing deployed model archive. "
        "The Rasa loader reads 496 examples across 21 intents from this pinned corpus; all classes meet the five-fold minimum. "
        "Native Rasa can drop classes with fewer than five examples in other corpora. "
        "Disclose actual coverage, eligible examples, fold count, source/config identity, warnings, "
        "accuracy and macro/weighted F1. Keep per-extractor entity scores distinct from the report's merged-span F1. "
        "Native shuffled splits do not establish grouped/paraphrase-independent generalisation. "
        "Never replace the existing held-out scores with cross-validation scores.")
    d.add_paragraph(
        "Verbatim source excerpt: rasa-bot/actions/actions.py, lines 434–451, "
        "published revision bbe39be9be5a50cd8ab48c2976201ff3b8d73806. "
        "The timeout prevents an indefinite provider wait; HTTP/JSON failures return None. "
        "This helper alone does not prove the user-facing fallback message or every failure branch was tested.")
    snippet = (
        "def _request_json(\n"
        "    method: str,\n"
        "    url: str,\n"
        "    *,\n"
        "    headers: dict[str, str] | None = None,\n"
        "    data: dict[str, Any] | None = None,\n"
        "    timeout: float = 2.3,\n"
        ") -> dict[str, Any] | None:\n"
        "    try:\n"
        "        response = requests.request(\n"
        "            method, url, headers=headers, json=data, timeout=timeout\n"
        "        )\n"
        "        response.raise_for_status()\n"
        "        return response.json()\n"
        "    except (requests.RequestException, ValueError) as exc:\n"
        "        logger.warning(\"External API request failed: %s\", exc)\n"
        "        return None"
    )
    code = d.add_paragraph()
    run = code.add_run(snippet)
    run.font.name = "Consolas"
    run.font.size = Pt(8)
    d.save(OUT / "Landa-User-Testing-Pack.docx")
    schemas = {
        "Task-Observations.csv": [
            "participant_id", "session_date", "app_version_or_revision", "device_browser",
            "task_id", "outcome", "elapsed_seconds", "time_capped", "assistance_count",
            "reached_recommendations", "anonymous_observation", "suggested_change"],
        "Survey-Responses.csv": [
            "participant_id", "Q1_clarity", "Q2_ease", "Q3_recommendation_explanation",
            "Q4_source_limitations", "Q5_control", "Q6_calibrated_reliance",
            "most_confusing", "most_helpful", "one_change"],
    }
    for filename, header in schemas.items():
        with (OUT / filename).open("w", newline="") as f:
            csv.writer(f).writerow(header)
    (OUT / "README.txt").write_text(
        "REAL-USER STUDY MATERIALS — no responses or scores are supplied.\n"
        "Read the Word/PDF protocol. Add rows to the CSV files only after voluntary real sessions.\n"
        "Use anonymous IDs and keep consent records locally. Do not upload identifying or private data.\n"
        "Task outcomes and survey responses are different measures; retain failed and blocked attempts.\n"
        "No public-study claim or grade prediction is supported by these blank forms.\n")
    print(OUT)


if __name__ == "__main__":
    main()